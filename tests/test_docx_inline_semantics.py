#!/usr/bin/env python3
"""Regression tests for hyperlinks, footnotes, and Word-native equations."""

from __future__ import annotations

import json
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import check_flow_b_gate
import import_docx
import omml_to_latex
import render_chapters
from docx import Document


def rewrite_docx_parts(
    docx_path: Path,
    replacements: dict[str, str],
    additions: dict[str, str] | None = None,
) -> None:
    """Replace selected OOXML parts without changing unrelated ZIP entries."""
    with tempfile.TemporaryDirectory() as tmp:
        rewritten = Path(tmp) / "rewritten.docx"
        with (
            zipfile.ZipFile(docx_path, "r") as source,
            zipfile.ZipFile(rewritten, "w") as target,
        ):
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in replacements:
                    data = replacements[info.filename].encode("utf-8")
                target.writestr(info, data)
            for filename, text in (additions or {}).items():
                target.writestr(filename, text.encode("utf-8"))
        docx_path.write_bytes(rewritten.read_bytes())


def test_import_and_render_supported_inline_docx_semantics():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "workspace/input").mkdir(parents=True)
        docx_path = root / "workspace/input/thesis.docx"
        document = Document()
        document.add_paragraph("INLINE_PLACEHOLDER")
        document.save(docx_path)

        with zipfile.ZipFile(docx_path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
            document_rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")

        inline_xml = (
            "<w:r><w:t>参见</w:t></w:r>"
            '<w:hyperlink r:id="rIdInlineLink">'
            "<w:r><w:rPr><w:b/></w:rPr><w:t>项目主页</w:t></w:r>"
            "</w:hyperlink>"
            "<w:r><w:t>并阅读说明</w:t></w:r>"
            '<w:r><w:footnoteReference w:id="1"/></w:r>'
            "<w:r><w:t>，模型为</w:t></w:r>"
            "<m:oMath><m:f>"
            "<m:num><m:r><m:t>a</m:t></m:r></m:num>"
            "<m:den><m:r><m:t>b</m:t></m:r></m:den>"
            "</m:f></m:oMath>"
            "<w:r><w:t>。</w:t></w:r>"
        )
        document_xml = document_xml.replace(
            "<w:r><w:t>INLINE_PLACEHOLDER</w:t></w:r>",
            inline_xml,
            1,
        )
        document_rels = document_rels.replace(
            "</Relationships>",
            '<Relationship Id="rIdInlineLink" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
            'Target="https://example.com/docs?x=1&amp;y=2" TargetMode="External"/>'
            "</Relationships>",
            1,
        )
        footnotes_xml = (
            '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:footnote w:id="1"><w:p>'
            "<w:r><w:footnoteRef/></w:r>"
            "<w:r><w:t>这是脚注内容。</w:t></w:r>"
            "</w:p></w:footnote>"
            "</w:footnotes>"
        )
        rewrite_docx_parts(
            docx_path,
            {
                "word/document.xml": document_xml,
                "word/_rels/document.xml.rels": document_rels,
            },
            {"word/footnotes.xml": footnotes_xml},
        )

        import_docx.extract(root, docx_path)
        thesis = json.loads(
            (root / "workspace/intermediate/thesis.json").read_text(encoding="utf-8")
        )
        block = next(block for block in thesis["source_blocks"] if block["id"] == "p0001")

        assert [item.get("kind", "text") for item in block["runs"]] == [
            "text",
            "hyperlink",
            "text",
            "footnote",
            "text",
            "equation",
            "text",
        ]
        assert thesis["counts"]["converted_inline_features"] == {
            "equations": 1,
            "footnotes": 1,
            "hyperlinks": 1,
        }
        assert thesis["unsupported_features"] == []
        assert check_flow_b_gate.block_inline_issues(block) == []

        latex = render_chapters.block_to_latex(block)
        assert r"\href{\detokenize{https://example.com/docs?x=1&y=2}}{\textbf{项目主页}}" in latex
        assert r"\footnote{这是脚注内容。}" in latex
        assert r"\(\frac{a}{b}\)" in latex


def test_omml_converter_reports_unknown_nodes_instead_of_claiming_success():
    math = ET.fromstring(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        "<m:unknown><m:r><m:t>x</m:t></m:r></m:unknown>"
        "</m:oMath>"
    )
    result = omml_to_latex.convert_omml(math)

    assert result.latex == "x"
    assert result.complete is False
    assert result.unsupported_tags == ("unknown",)


def test_table_equations_convert_but_table_footnotes_remain_explicit_boundary():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "workspace/input").mkdir(parents=True)
        docx_path = root / "workspace/input/thesis.docx"
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "TABLE_EQUATION"
        table.cell(0, 1).text = "TABLE_FOOTNOTE"
        document.save(docx_path)

        with zipfile.ZipFile(docx_path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        document_xml = document_xml.replace(
            "<w:r><w:t>TABLE_EQUATION</w:t></w:r>",
            "<m:oMath><m:sSup>"
            "<m:e><m:r><m:t>x</m:t></m:r></m:e>"
            "<m:sup><m:r><m:t>2</m:t></m:r></m:sup>"
            "</m:sSup></m:oMath>",
            1,
        ).replace(
            "<w:r><w:t>TABLE_FOOTNOTE</w:t></w:r>",
            '<w:r><w:t>单元格</w:t><w:footnoteReference w:id="2"/></w:r>',
            1,
        )
        footnotes_xml = (
            '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:footnote w:id="2"><w:p><w:r><w:t>表格脚注</w:t></w:r></w:p></w:footnote>'
            "</w:footnotes>"
        )
        rewrite_docx_parts(
            docx_path,
            {"word/document.xml": document_xml},
            {"word/footnotes.xml": footnotes_xml},
        )

        import_docx.extract(root, docx_path)
        thesis = json.loads(
            (root / "workspace/intermediate/thesis.json").read_text(encoding="utf-8")
        )
        table_block = next(
            block for block in thesis["source_blocks"] if block["source_type"] == "table"
        )
        equation = table_block["table"]["inline_rows"][0][0][0][0]
        footnote = table_block["table"]["inline_rows"][0][1][0][1]

        assert equation["kind"] == "equation"
        assert equation["latex"] == "x^{2}"
        assert footnote["kind"] == "footnote"
        assert footnote["conversion_status"] == "needs_confirmation"
        assert {feature["type"] for feature in thesis["unsupported_features"]} == {
            "unconverted_footnote"
        }
        assert any(
            issue["check"] == "inline_conversion_status"
            for issue in check_flow_b_gate.block_inline_issues(table_block)
        )


def test_omml_converter_handles_common_empirical_formula_structure():
    math = ET.fromstring(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:nary><m:naryPr><m:chr m:val="∑"/></m:naryPr>'
        "<m:sub><m:r><m:t>i=1</m:t></m:r></m:sub>"
        "<m:sup><m:r><m:t>n</m:t></m:r></m:sup>"
        "<m:e><m:sSubSup>"
        "<m:e><m:r><m:t>x</m:t></m:r></m:e>"
        "<m:sub><m:r><m:t>i</m:t></m:r></m:sub>"
        "<m:sup><m:r><m:t>2</m:t></m:r></m:sup>"
        "</m:sSubSup></m:e></m:nary>"
        "</m:oMath>"
    )
    result = omml_to_latex.convert_omml(math)

    assert result.complete is True
    assert result.latex == r"\sum_{i=1}^{n} x_{i}^{2}"


def test_omml_converter_preserves_run_boundaries_and_no_bar_fractions():
    math = ET.fromstring(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        "<m:r><m:t>α</m:t></m:r><m:r><m:t>x</m:t></m:r>"
        '<m:f><m:fPr><m:type m:val="noBar"/></m:fPr>'
        "<m:num><m:r><m:t>n</m:t></m:r></m:num>"
        "<m:den><m:r><m:t>k</m:t></m:r></m:den>"
        "</m:f>"
        "</m:oMath>"
    )
    result = omml_to_latex.convert_omml(math)

    assert result.complete is True
    assert result.latex == r"\alpha{}x\genfrac{}{}{0pt}{}{n}{k}"


def test_omml_converter_preserves_supported_math_script_properties():
    math = ET.fromstring(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:r><m:rPr><m:scr m:val="double-struck"/></m:rPr><m:t>R</m:t></m:r>'
        "</m:oMath>"
    )
    result = omml_to_latex.convert_omml(math)

    assert result.complete is True
    assert result.latex == r"\mathbb{R}"
