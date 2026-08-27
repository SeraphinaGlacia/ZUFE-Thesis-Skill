#!/usr/bin/env python3
"""Focused regression tests extracted from the former monolithic suite."""

from __future__ import annotations

import json
import tempfile
from pathlib import Path

from docx import Document

from .support import (
    basic_metadata_yaml,
    load_module,
)


def test_metadata_parser_preserves_numeric_identifiers_with_leading_zeroes():
    common = load_module("common")
    with tempfile.TemporaryDirectory() as tmp:
        metadata_path = Path(tmp) / "metadata.yaml"
        metadata_path.write_text(
            "student_id: 00123\nclass_name: 0007\nreport_style: 1\n",
            encoding="utf-8",
        )
        metadata = common.load_metadata_yaml(metadata_path)

        assert metadata["student_id"] == "00123"
        assert metadata["class_name"] == "0007"
        assert metadata["report_style"] == "1"


def test_prescan_reads_cover_table_metadata_without_report_style_default():
    prescan_docx = load_module("prescan_docx")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        docx_path = root / "cover.docx"
        document = Document()
        document.add_paragraph("专业实践报告")
        table = document.add_table(rows=4, cols=2)
        rows = [
            ("指导教师", "张老师"),
            ("专业名称", "数字经济"),
            ("学院", "经济学院"),
            ("日期", "2026年6月"),
        ]
        for row, (label, value) in zip(table.rows, rows, strict=True):
            row.cells[0].text = label
            row.cells[1].text = value
        document.save(docx_path)

        result = prescan_docx.prescan(root, docx_path)
        candidates = result["metadata_candidates"]
        assert candidates["report_style"] == "1"
        assert candidates["mentor"] == "张老师"
        assert candidates["major"] == "数字经济"
        assert candidates["college"] == "经济学院"
        assert candidates["date"] == "2026年6月"

        assert prescan_docx.metadata_candidates(["普通论文标题"])["report_style"] == ""


def test_prescan_and_import_accept_table_only_docx_metadata():
    import_docx = load_module("import_docx")
    prescan_docx = load_module("prescan_docx")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "workspace/input").mkdir(parents=True)
        docx_path = root / "workspace/input/thesis.docx"
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "学生姓名"
        table.cell(0, 1).text = "张三"
        table.cell(1, 0).text = "学号"
        table.cell(1, 1).text = "20260001"
        document.save(docx_path)

        prescan_result = prescan_docx.prescan(root, docx_path)
        import_result = import_docx.extract(root, docx_path)
        thesis = json.loads(
            (root / "workspace/intermediate/thesis.json").read_text(encoding="utf-8")
        )

        assert prescan_result["status"] == "passed"
        assert prescan_result["counts"]["non_empty_paragraphs"] == 0
        assert prescan_result["metadata_candidates"]["name"] == "张三"
        assert import_result["counts"]["tables"] == 1
        assert thesis["metadata_candidates"]["name"] == "张三"
        assert thesis["metadata_candidates"]["student_id"] == "20260001"


def test_render_basicinfo_blocks_missing_report_style():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        metadata = root / "metadata.yaml"
        metadata.write_text("thesis_title_cn: 测试题目\n", encoding="utf-8")

        result = render_basicinfo.render(root, metadata, thesis_path=None)
        assert result["status"] == "blocked"
        assert "report_style" in result["missing_fields"]
        assert not (root / "chapters/basicinfo.tex").exists()


def test_render_basicinfo_blocks_missing_required_cover_metadata():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        metadata = root / "metadata.yaml"
        metadata.write_text(
            "report_style: 1\nenglish_content_decision: omit\n",
            encoding="utf-8",
        )

        result = render_basicinfo.render(root, metadata, thesis_path=None)
        assert result["status"] == "blocked"
        assert result["gate"] == "metadata_required"
        assert "thesis_title_cn" in result["missing_fields"]
        assert "name" in result["missing_fields"]
        assert "student_id" in result["missing_fields"]
        assert not (root / "chapters/basicinfo.tex").exists()


def test_render_basicinfo_blocks_missing_subtitle_when_enabled():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        metadata = root / "metadata.yaml"
        metadata.write_text(
            basic_metadata_yaml(
                extra="has_subtitle: true\nenglish_content_decision: omit\n",
            ),
            encoding="utf-8",
        )

        result = render_basicinfo.render(root, metadata, thesis_path=None)
        assert result["status"] == "blocked"
        assert result["gate"] == "metadata_required"
        assert "thesis_subtitle_cn" in result["missing_fields"]
        assert "thesis_subtitle_en" in result["missing_fields"]
        assert not (root / "chapters/basicinfo.tex").exists()


def test_render_basicinfo_blocks_unapproved_generated_english():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        (root / "workspace/intermediate").mkdir(parents=True)
        metadata = root / "metadata.yaml"
        metadata.write_text(basic_metadata_yaml(), encoding="utf-8")
        thesis_path = root / "workspace/intermediate/thesis.json"
        thesis_path.write_text(
            json.dumps(
                {
                    "metadata": {
                        "abstract_en": "Generated English abstract.",
                        "keywords_en": ["generated", "keywords"],
                        "english_content_source": "generated",
                    }
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        result = render_basicinfo.render(root, metadata, thesis_path)
        assert result["status"] == "blocked"
        assert result["gate"] == "generated_english_requires_confirmation"
        assert not (root / "chapters/basicinfo.tex").exists()

        metadata.write_text(
            basic_metadata_yaml(extra="allow_generated_english: true\n"),
            encoding="utf-8",
        )
        result = render_basicinfo.render(root, metadata, thesis_path)
        assert result["status"] == "passed"


def test_render_basicinfo_requires_missing_english_content_decision():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        (root / "workspace/intermediate").mkdir(parents=True)
        metadata = root / "metadata.yaml"
        metadata.write_text(basic_metadata_yaml(), encoding="utf-8")
        thesis_path = root / "workspace/intermediate/thesis.json"
        thesis_path.write_text(
            json.dumps({"metadata": {"abstract_cn": "中文摘要。"}}, ensure_ascii=False),
            encoding="utf-8",
        )

        result = render_basicinfo.render(root, metadata, thesis_path)
        assert result["status"] == "blocked"
        assert result["gate"] == "english_content_decision_required"
        assert not (root / "chapters/basicinfo.tex").exists()

        metadata.write_text(
            basic_metadata_yaml(extra="english_content_decision: omit\n"),
            encoding="utf-8",
        )
        result = render_basicinfo.render(root, metadata, thesis_path)
        assert result["status"] == "passed"


def test_render_basicinfo_requires_explicit_source_field_evidence():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        (root / "workspace/intermediate").mkdir(parents=True)
        metadata = root / "metadata.yaml"
        metadata.write_text(
            basic_metadata_yaml(extra="english_content_decision: omit\n"),
            encoding="utf-8",
        )
        thesis_path = root / "workspace/intermediate/thesis.json"
        thesis_path.write_text(
            json.dumps(
                {
                    "metadata": {},
                    "source_blocks": [
                        {
                            "id": "p0001",
                            "status": "mapped",
                            "source_type": "paragraph",
                            "text": "绝不能丢失的封面附注",
                            "target_slot": "chapters/basicinfo.tex",
                            "render_result": None,
                        }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        result = render_basicinfo.render(root, metadata, thesis_path)
        thesis = json.loads(thesis_path.read_text(encoding="utf-8"))
        assert result["status"] == "needs_confirmation"
        assert result["gate"] == "basicinfo_source_evidence"
        assert result["unverified_blocks"][0]["reasons"] == ["missing_metadata_fields"]
        assert thesis["source_blocks"][0]["status"] == "mapped"
        assert thesis["source_blocks"][0]["render_result"] is None
        assert "绝不能丢失的封面附注" not in (root / "chapters/basicinfo.tex").read_text(
            encoding="utf-8"
        )


def test_render_basicinfo_marks_only_verified_metadata_bindings_rendered():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        (root / "workspace/intermediate").mkdir(parents=True)
        metadata = root / "metadata.yaml"
        metadata.write_text(
            basic_metadata_yaml(extra="english_content_decision: omit\n"),
            encoding="utf-8",
        )
        thesis_path = root / "workspace/intermediate/thesis.json"
        thesis_path.write_text(
            json.dumps(
                {
                    "metadata": {},
                    "source_blocks": [
                        {
                            "id": "p0001",
                            "status": "mapped",
                            "source_type": "paragraph",
                            "text": "学生姓名：测试姓名",
                            "metadata_fields": ["name"],
                            "target_slot": "chapters/basicinfo.tex",
                            "render_result": None,
                        }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        result = render_basicinfo.render(root, metadata, thesis_path)
        thesis = json.loads(thesis_path.read_text(encoding="utf-8"))
        block = thesis["source_blocks"][0]
        assert result["status"] == "passed"
        assert block["status"] == "rendered"
        assert block["render_result"]["metadata_fields"] == ["name"]
        assert block["render_result"]["evidence"] == "all_bound_values_found_in_source_block"


def test_render_basicinfo_requires_an_explicit_destination_for_residual_text():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        (root / "workspace/intermediate").mkdir(parents=True)
        metadata = root / "metadata.yaml"
        metadata.write_text(
            basic_metadata_yaml(extra="english_content_decision: omit\n"),
            encoding="utf-8",
        )
        thesis_path = root / "workspace/intermediate/thesis.json"
        thesis_path.write_text(
            json.dumps(
                {
                    "metadata": {},
                    "source_blocks": [
                        {
                            "id": "p0001",
                            "status": "mapped",
                            "source_type": "paragraph",
                            "text": "学生姓名：测试姓名；附注：不得公开",
                            "metadata_fields": ["name"],
                            "target_slot": "chapters/basicinfo.tex",
                            "render_result": None,
                        }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        result = render_basicinfo.render(root, metadata, thesis_path)
        assert result["status"] == "needs_confirmation"
        assert any(
            reason.startswith("uncovered_source_text:")
            for reason in result["unverified_blocks"][0]["reasons"]
        )

        thesis = json.loads(thesis_path.read_text(encoding="utf-8"))
        block = thesis["source_blocks"][0]
        block["metadata_excluded_text"] = ["附注：不得公开"]
        block["metadata_exclusion_reason"] = "隐私提示，不属于论文封面输出。"
        thesis_path.write_text(json.dumps(thesis, ensure_ascii=False), encoding="utf-8")

        result = render_basicinfo.render(root, metadata, thesis_path)
        rendered = json.loads(thesis_path.read_text(encoding="utf-8"))["source_blocks"][0]
        assert result["status"] == "passed"
        assert rendered["status"] == "rendered"


def test_render_basicinfo_accepts_report_style_candidate_synonyms():
    render_basicinfo = load_module("render_basicinfo")
    verified, problems = render_basicinfo.verify_basicinfo_block(
        {
            "text": "实践报告",
            "metadata_fields": ["report_style"],
        },
        {"report_style": "1"},
    )
    assert verified == ["report_style"]
    assert problems == []


def test_render_basicinfo_blocks_malformed_ledger_without_writing():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        (root / "workspace/intermediate").mkdir(parents=True)
        metadata = root / "metadata.yaml"
        metadata.write_text(
            basic_metadata_yaml(extra="english_content_decision: omit\n"),
            encoding="utf-8",
        )
        thesis_path = root / "workspace/intermediate/thesis.json"
        thesis_path.write_text(
            json.dumps({"metadata": {}, "source_blocks": {}}, ensure_ascii=False),
            encoding="utf-8",
        )

        result = render_basicinfo.render(root, metadata, thesis_path)
        assert result["status"] == "blocked"
        assert result["gate"] == "thesis_json_structure"
        assert not (root / "chapters/basicinfo.tex").exists()


def test_render_basicinfo_supports_thesis_title_abs():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        metadata = root / "metadata.yaml"
        metadata.write_text(
            basic_metadata_yaml(
                title_cn="封面题目",
                extra="thesis_title_abs_cn: 摘要页题目\nenglish_content_decision: omit\n",
            ),
            encoding="utf-8",
        )
        render_basicinfo.render(root, metadata, thesis_path=None)
        basicinfo = (root / "chapters/basicinfo.tex").read_text(encoding="utf-8")
        assert "\\newcommand{\\thesisTitle}{封面题目}" in basicinfo
        assert "\\newcommand{\\thesisTitleAbs}{摘要页题目}" in basicinfo


def test_render_basicinfo_hides_hyperref_link_borders():
    render_basicinfo = load_module("render_basicinfo")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "chapters").mkdir()
        metadata = root / "metadata.yaml"
        metadata.write_text(
            basic_metadata_yaml(extra="english_content_decision: omit\n"),
            encoding="utf-8",
        )
        render_basicinfo.render(root, metadata, thesis_path=None)
        basicinfo = (root / "chapters/basicinfo.tex").read_text(encoding="utf-8")
        assert r"\hypersetup{hidelinks,pdfborder={0 0 0},pdfborderstyle={/S/U/W 0}}" in basicinfo
