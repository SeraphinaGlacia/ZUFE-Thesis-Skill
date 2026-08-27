#!/usr/bin/env python3
"""Focused regression tests extracted from the former monolithic suite."""

from __future__ import annotations

import json
import tempfile
import zipfile
from pathlib import Path

from docx import Document

from .support import (
    load_module,
    rewrite_docx_xml,
    write_tiny_png,
)


def test_import_docx_preserves_superscript_runs():
    import_docx = load_module("import_docx")
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("引用")
    ref = paragraph.add_run("1")
    ref.font.superscript = True

    runs = import_docx.run_payload(paragraph)
    assert runs == [
        {
            "index": 1,
            "text": "引用",
            "bold": False,
            "italic": False,
            "superscript": False,
            "subscript": False,
            "font_size_pt": None,
        },
        {
            "index": 2,
            "text": "1",
            "bold": False,
            "italic": False,
            "superscript": True,
            "subscript": False,
            "font_size_pt": None,
        },
    ]


def test_import_docx_preserves_image_anchor_order():
    import_docx = load_module("import_docx")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "workspace/input").mkdir(parents=True)
        png = root / "anchor.png"
        write_tiny_png(png)
        document = Document()
        document.add_paragraph("图片前段落")
        document.add_picture(str(png))
        document.add_paragraph("图片后段落")
        docx_path = root / "workspace/input/thesis.docx"
        document.save(docx_path)

        import_docx.extract(root, docx_path)
        thesis = json.loads(
            (root / "workspace/intermediate/thesis.json").read_text(encoding="utf-8")
        )
        blocks = thesis["source_blocks"]
        before = next(block for block in blocks if block.get("text") == "图片前段落")
        after = next(block for block in blocks if block.get("text") == "图片后段落")
        image = next(block for block in blocks if block.get("source_type") == "image")

        assert before["order"] < image["order"] < after["order"]
        assert image["status"] == "needs_confirmation"
        assert image["asset_status"] == "pending_export"
        assert image["target_slot"] is None
        assert image["evidence"]["docx_media_path"].startswith("word/media/")
        assert image["evidence"]["anchor_paragraph_id"] == "p0002"
        assert thesis["counts"]["paragraphs"] == 3
        assert thesis["counts"]["source_blocks_by_type"]["paragraph"] == 2
        assert thesis["counts"]["source_blocks_by_type"]["image"] == 1


def test_import_docx_preserves_repeated_uses_of_the_same_image_relationship():
    import_docx = load_module("import_docx")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "workspace/input").mkdir(parents=True)
        png = root / "repeated.png"
        write_tiny_png(png)
        document = Document()
        paragraph = document.add_paragraph("重复图片")
        run = paragraph.add_run()
        run.add_picture(str(png))
        run.add_picture(str(png))
        docx_path = root / "workspace/input/thesis.docx"
        document.save(docx_path)

        import_docx.extract(root, docx_path)
        thesis = json.loads(
            (root / "workspace/intermediate/thesis.json").read_text(encoding="utf-8")
        )
        images = [block for block in thesis["source_blocks"] if block.get("source_type") == "image"]

        assert len(images) == 2
        assert images[0]["evidence"]["docx_media_path"] == images[1]["evidence"]["docx_media_path"]
        assert [image["evidence"]["anchor_image_occurrence"] for image in images] == [1, 2]


def test_export_assets_does_not_mark_image_semantic_position_mapped():
    import_docx = load_module("import_docx")
    export_assets = load_module("export_assets")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "workspace/input").mkdir(parents=True)
        png = root / "anchor.png"
        write_tiny_png(png)
        document = Document()
        document.add_paragraph("图片前段落")
        document.add_picture(str(png))
        document.add_paragraph("图片后段落")
        docx_path = root / "workspace/input/thesis.docx"
        document.save(docx_path)

        import_docx.extract(root, docx_path)
        thesis_path = root / "workspace/intermediate/thesis.json"
        export_assets.export_assets(root, docx_path, thesis_path)
        thesis = json.loads(thesis_path.read_text(encoding="utf-8"))
        image = next(
            block for block in thesis["source_blocks"] if block.get("source_type") == "image"
        )

        assert image["status"] == "needs_confirmation"
        assert image["target_slot"] is None
        assert image["asset_status"] == "exported"
        assert image["asset_output"].startswith("Images/word_media/")
        assert image["render_result"]["kind"] == "asset_extracted"


def test_import_docx_reports_unsupported_features():
    import_docx = load_module("import_docx")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "workspace/input").mkdir(parents=True)
        document = Document()
        document.add_paragraph("正文段落")
        docx_path = root / "workspace/input/thesis.docx"
        document.save(docx_path)

        with zipfile.ZipFile(docx_path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        insertion = (
            "<w:p><w:hyperlink><w:r><w:t>链接文本</w:t></w:r></w:hyperlink></w:p>"
            "<w:p><w:r><m:oMath><m:r><m:t>x=1</m:t></m:r></m:oMath></w:r></w:p>"
            '<w:p><w:ins w:id="1" w:author="tester"><w:r><w:t>修订文本</w:t></w:r></w:ins></w:p>'
            "<w:p><w:r><w:pict><v:textbox><w:txbxContent><w:p><w:r><w:t>文本框</w:t></w:r></w:p></w:txbxContent></v:textbox></w:pict></w:r></w:p>"
            "<w:p><w:pPr><w:numPr/></w:pPr><w:r><w:t>自动编号</w:t></w:r></w:p>"
            '<w:p><w:fldSimple w:instr="REF target"><w:r><w:t>域结果</w:t></w:r></w:fldSimple></w:p>'
            '<w:p><w:r><w:drawing><w:blip r:link="rId999"/></w:drawing></w:r></w:p>'
            "<w:p><w:r><w:drawing><chart/></w:drawing></w:r></w:p>"
            "<w:p><w:r><w:object><OLEObject/></w:object></w:r></w:p>"
        )
        rewrite_docx_xml(
            docx_path,
            {"word/document.xml": document_xml.replace("<w:sectPr", insertion + "<w:sectPr", 1)},
            {
                "word/footnotes.xml": (
                    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:footnote w:id="1"><w:p><w:r><w:t>脚注</w:t></w:r></w:p></w:footnote>'
                    "</w:footnotes>"
                ),
                "word/comments.xml": (
                    '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:comment w:id="1"><w:p><w:r><w:t>批注</w:t></w:r></w:p></w:comment>'
                    "</w:comments>"
                ),
                "word/header1.xml": (
                    '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    "<w:p><w:r><w:t>页眉</w:t></w:r></w:p></w:hdr>"
                ),
            },
        )

        import_docx.extract(root, docx_path)
        thesis = json.loads(
            (root / "workspace/intermediate/thesis.json").read_text(encoding="utf-8")
        )
        features = {feature["type"]: feature for feature in thesis["unsupported_features"]}

        assert features["hyperlink"]["count"] == 1
        assert features["equation_omml"]["count"] == 1
        assert features["tracked_changes"]["count"] == 1
        assert features["textbox"]["count"] == 1
        assert features["footnote_or_endnote"]["count"] == 1
        assert features["comment"]["count"] == 1
        assert features["header_footer"]["count"] == 1
        assert features["linked_image"]["count"] == 1
        assert features["chart_or_smartart"]["count"] == 1
        assert features["ole_object"]["count"] >= 1
        assert features["field_code"]["count"] >= 1
        assert features["automatic_numbering"]["count"] == 1
        for feature in features.values():
            assert feature["status"] == "needs_confirmation"
            assert feature["locations"]


def test_import_and_prescan_expand_block_content_controls():
    import_docx = load_module("import_docx")
    prescan_docx = load_module("prescan_docx")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "workspace/input").mkdir(parents=True)
        document = Document()
        document.add_paragraph("普通正文")
        docx_path = root / "workspace/input/thesis.docx"
        document.save(docx_path)

        with zipfile.ZipFile(docx_path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        content_control = (
            '<w:sdt><w:sdtPr><w:tag w:val="audit"/></w:sdtPr><w:sdtContent>'
            "<w:p><w:r><w:t>内容控件正文</w:t></w:r></w:p>"
            "</w:sdtContent></w:sdt>"
        )
        rewrite_docx_xml(
            docx_path,
            {
                "word/document.xml": document_xml.replace(
                    "<w:sectPr", content_control + "<w:sectPr", 1
                )
            },
        )

        prescan_result = prescan_docx.prescan(root, docx_path)
        assert prescan_result["status"] == "passed"
        assert prescan_result["counts"]["non_empty_paragraphs"] == 2
        assert any(block["text"] == "内容控件正文" for block in prescan_result["structure_preview"])

        import_docx.extract(root, docx_path)
        thesis = json.loads(
            (root / "workspace/intermediate/thesis.json").read_text(encoding="utf-8")
        )
        controlled = next(
            block for block in thesis["source_blocks"] if block.get("text") == "内容控件正文"
        )
        features = {feature["type"]: feature for feature in thesis["unsupported_features"]}
        assert controlled["evidence"]["inside_content_control"] is True
        assert features["content_control"]["count"] == 1


def test_export_assets_blocks_changed_source_docx():
    import_docx = load_module("import_docx")
    export_assets = load_module("export_assets")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "workspace/input").mkdir(parents=True)
        document = Document()
        document.add_paragraph("正文")
        docx_path = root / "workspace/input/thesis.docx"
        document.save(docx_path)

        import_docx.extract(root, docx_path)
        thesis_path = root / "workspace/intermediate/thesis.json"
        original_thesis = thesis_path.read_text(encoding="utf-8")
        docx_path.write_bytes(docx_path.read_bytes() + b"changed-after-import")

        result = export_assets.export_assets(root, docx_path, thesis_path)
        assert result["status"] == "blocked"
        assert result["gate"] == "source_docx_changed"
        assert not (root / "Images/word_media").exists()
        assert thesis_path.read_text(encoding="utf-8") == original_thesis
